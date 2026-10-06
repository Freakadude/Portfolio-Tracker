"""What kind of file a holdings download is, and Excel workbooks read into plain tables (FR-MD-09).

The kind is decided by the file's first bytes, never by its name: an issuer's download link often
has no extension. Every reader here ends in the same thing, a list of rows of text, so the finding
of the header, the matching of the columns and the checks of the weights stay in `parse.py`.
"""

from __future__ import annotations

import io
import zipfile
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from typing import Literal

from folio.imports.parse import ParseError

HOLDINGS_MAX_BYTES = 10 * 1024 * 1024  # workbooks and PDFs are larger than a CSV of the same list
MAX_UNPACKED_BYTES = 100 * 1024 * 1024  # what a workbook may unpack to (a zip bomb stops here)
MAX_SHEET_ROWS = 50_000

FileKind = Literal["xlsx", "pdf", "xls", "csv"]
Table = list[list[str]]

_OLE2 = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # the old binary Excel format (.xls)


def detect_kind(data: bytes) -> FileKind:
    head = data[:8]
    if head.startswith(b"PK\x03\x04"):
        return "xlsx"
    if data[:1024].lstrip().startswith(b"%PDF-"):
        return "pdf"
    if head == _OLE2:
        return "xls"
    return "csv"


def _text(value: object) -> str:
    """A cell as text. Numbers keep the digits they were typed with (7.12 stays 7.12)."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        try:
            return format(Decimal(repr(value)), "f")
        except InvalidOperation:
            return str(value)
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat(" ")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def read_workbook(data: bytes) -> dict[str, Table]:
    """Every visible sheet of an .xlsx workbook as rows of text, empty rows left out. Only the
    values Excel last stored are read: no formula is evaluated and no macro is run."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = set(archive.namelist())
            unpacked = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise ParseError("The file is not a readable Excel workbook (.xlsx).") from exc
    if "xl/workbook.xml" not in names:
        raise ParseError("That file is a zip archive, not an Excel workbook (.xlsx).")
    if unpacked > MAX_UNPACKED_BYTES:
        raise ParseError("The workbook is too large when unpacked (over 100 MB).")

    from openpyxl import load_workbook  # imported when needed: the CSV path never loads it

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many kinds for a damaged file
        raise ParseError("The Excel workbook could not be read. Is it damaged?") from exc
    tables: dict[str, Table] = {}
    try:
        for sheet in workbook.worksheets:
            if sheet.sheet_state != "visible":
                continue
            sheet.reset_dimensions()  # some issuers' files state a one-cell size
            rows: Table = []
            for row in sheet.iter_rows(values_only=True, max_row=MAX_SHEET_ROWS):
                cells = [_text(v) for v in row]
                if any(cells):
                    rows.append(cells)
            if rows:
                tables[sheet.title] = rows
    finally:
        workbook.close()
    if not tables:
        raise ParseError("The workbook has no sheet with any content.")
    return tables
