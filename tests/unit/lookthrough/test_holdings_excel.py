"""Holdings from Excel workbooks (FR-MD-09): the layouts issuers use, built here in the test.
The expected holdings, weights and coverage are counted by hand."""

import io
import zipfile
from datetime import date
from decimal import Decimal
from typing import Any

import pytest
from openpyxl import Workbook

from folio.imports.parse import ParseError
from folio.lookthrough.formats import detect_kind
from folio.lookthrough.parse import HoldingsMapping, read_holdings, suggest

D = Decimal


def workbook(sheets: dict[str, list[list[Any]]]) -> bytes:
    book = Workbook()
    first = True
    for name, rows in sheets.items():
        sheet = book.active if first else book.create_sheet()
        first = False
        sheet.title = name
        for row in rows:
            sheet.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


ENGLISH = [
    ["Example Core ETF"],
    ["Holdings as of", date(2026, 10, 2)],
    [],
    ["Name", "ISIN", "Sector", "Country", "Weight (%)"],
    ["ALPHA TECH INC", "US0000000001", "Technology", "United States", 60.5],
    ["BETA BANK PLC", "GB0000000002", "Financials", "United Kingdom", 30.25],
    ["CASH", None, "Cash", "United States", 0.25],
    ["CHIP MAKER NV", "NL0000000003", "Technology", "Netherlands", 9.0],
    [],
    ["Source: the issuer. Not an offer."],
]


def test_a_workbook_is_recognised_by_its_content_not_its_name() -> None:
    assert detect_kind(workbook({"A": [["x"]]})) == "xlsx"
    assert detect_kind(b"%PDF-1.7\n...") == "pdf"
    assert detect_kind(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1rest") == "xls"
    assert detect_kind(b"Name;Weight\nA;1\n") == "csv"


def test_a_holdings_sheet_is_read_with_its_date_and_cash_left_out() -> None:
    file = suggest(workbook({"Holdings": ENGLISH}))
    assert file.as_of == date(2026, 10, 2)
    assert file.mapping.sheet is None or file.mapping.sheet == "Holdings"
    assert file.sheets == []  # one sheet: nothing to choose
    read = read_holdings(file)
    assert [(c.name, c.weight_pct) for c in read.constituents] == [
        ("ALPHA TECH INC", D("60.5")),
        ("BETA BANK PLC", D("30.25")),
        ("CHIP MAKER NV", D("9.0")),
    ]
    assert read.covered_pct == D("99.75")
    assert read.constituents[0].isin == "US0000000001"
    assert read.constituents[0].country == "United States"
    assert len(read.dropped) == 1 and read.dropped[0].startswith("CASH")
    assert read.errors == []


def test_the_holdings_sheet_is_found_behind_a_cover_sheet() -> None:
    cover = [["Fund factsheet"], ["Please read the disclaimer"], ["Page 1"]]
    file = suggest(workbook({"Cover": cover, "Holdings": ENGLISH, "Notes": [["a", "b"]]}))
    assert file.mapping.sheet == "Holdings"
    info = {s.name: s.holdings for s in file.sheets}
    assert info == {"Cover": 0, "Holdings": 4, "Notes": 0}
    assert len(read_holdings(file).constituents) == 3


def test_columns_in_another_order_and_in_german_are_matched_by_name() -> None:
    german = [
        ["Fonds XY", None, None],
        ["Stand:", "02.10.2026", None],
        ["Gewicht (%)", "Name", "Land"],
        [50, "ALPHA AG", "Deutschland"],
        [50, "BETA SA", "Frankreich"],
    ]
    read = read_holdings(suggest(workbook({"Bestand": german})))
    assert [(c.name, c.weight_pct, c.country) for c in read.constituents] == [
        ("ALPHA AG", D(50), "Deutschland"),
        ("BETA SA", D(50), "Frankreich"),
    ]
    assert read.covered_pct == D(100)


def test_percent_formatted_cells_stored_as_fractions_are_read_as_percent() -> None:
    rows = [["Name", "Weighting"], ["A", 0.6], ["B", 0.4]]
    file = suggest(workbook({"S": rows}))
    assert file.mapping.weight_is_fraction is True
    read = read_holdings(file)
    assert [c.weight_pct for c in read.constituents] == [D(60), D(40)]
    assert read.covered_pct == D(100)


def test_numbers_keep_the_digits_they_were_typed_with() -> None:
    rows = [["Name", "Weight (%)"], ["A", 7.12], ["B", 92.88]]
    read = read_holdings(suggest(workbook({"S": rows})))
    assert [c.weight_pct for c in read.constituents] == [D("7.12"), D("92.88")]


def test_another_sheet_can_be_chosen_and_is_read() -> None:
    other = [["Name", "Weight (%)"], ["X", 100]]
    data = workbook({"Holdings": ENGLISH, "Other": other})
    assert suggest(data).mapping.sheet == "Holdings"  # the one with more holdings
    chosen = suggest(data, "Other")
    assert chosen.mapping.sheet == "Other"
    assert [c.name for c in read_holdings(chosen).constituents] == ["X"]
    with pytest.raises(ParseError, match="no sheet called 'Missing'"):
        suggest(data, "Missing")


def test_a_hidden_sheet_is_left_out() -> None:
    book = Workbook()
    shown = book.active
    shown.title = "Holdings"
    for row in ENGLISH:
        shown.append(row)
    hidden = book.create_sheet("Old")
    hidden.sheet_state = "hidden"
    hidden.append(["Name", "Weight (%)"])
    hidden.append(["Z", 100])
    out = io.BytesIO()
    book.save(out)
    file = suggest(out.getvalue())
    assert file.sheets == [] and file.mapping.sheet == "Holdings"


def test_a_workbook_without_a_holdings_table_says_so() -> None:
    with pytest.raises(ParseError, match="on any sheet"):
        suggest(workbook({"A": [["hello", "world"]], "B": [["x", "y"]]}))


def test_files_that_cannot_be_read_get_a_plain_message() -> None:
    with pytest.raises(ParseError, match="old Excel file"):
        suggest(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"0" * 100)
    with pytest.raises(ParseError, match="not a readable Excel workbook"):
        suggest(b"PK\x03\x04 this is not a zip file")
    plain_zip = io.BytesIO()
    with zipfile.ZipFile(plain_zip, "w") as z:
        z.writestr("readme.txt", "hello")
    with pytest.raises(ParseError, match="zip archive, not an Excel workbook"):
        suggest(plain_zip.getvalue())
    with pytest.raises(ParseError, match="larger than 10 MB"):
        suggest(b"a," * (6 * 1024 * 1024))
    with pytest.raises(ParseError, match="PDF could not be read"):
        suggest(b"%PDF-1.7\n1 0 obj")


def test_a_workbook_that_unpacks_to_too_much_is_refused(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("folio.lookthrough.formats.MAX_UNPACKED_BYTES", 10)
    with pytest.raises(ParseError, match="too large when unpacked"):
        suggest(workbook({"S": ENGLISH}))


def test_a_mapping_given_by_the_owner_is_used_on_the_chosen_sheet() -> None:
    data = workbook({"Holdings": ENGLISH, "Other": [["Name", "Weight (%)"], ["X", 100]]})
    file = suggest(data, "Holdings")
    mapping = HoldingsMapping.model_validate(
        {**file.mapping.model_dump(), "sector": 3}  # the country column, as a sector
    )
    read = read_holdings(file, mapping)
    assert read.constituents[0].sector == "United States"


GERMAN_ISHARES = [
    ["All"],
    ["Per", "05.Okt.2026"],
    [
        "Emittententicker", "Name", "Sektor", "Anlageklasse", "Marktwert", "Gewichtung (%)",
        "Nominalwert", "Nominale", "Marktw\u00e4hrung",
    ],
    ["ASML", "ASML HOLDING", "IT", "Aktien", "EUR 837.207.444", "5,32", "837.207.443,80", "5", "EUR"],
    ["HSBA", "HSBC HOLDINGS PLC", "Finanzwesen", "Aktien", "EUR 385.135.760", "94,25", "1,0", "2", "GBP"],
    ["EUR", "EUR CASH", "Barmittel & Derivate", "Geldmarkt", "EUR 36.025.636", "0,23", "3", "4", "EUR"],
    ["ZRPZ6", "MSCI EUROPE INDEX DEC 26", "Barmittel & Derivate", "Futures", "EUR 0", "0,00", "1", "2", "EUR"],
    ["GBP", "GBP/EUR", "Barmittel & Derivate", "FX", "EUR -76.520", "0,00", "1", "2", "EUR"],
    ["\u00bbFondspositionen und Kennzahlen\u00ab enth\u00e4lt eine detaillierte Aufstellung."],
]  # fmt: skip


def test_an_issuers_german_workbook_is_read_with_its_cash_lines_and_date() -> None:
    """The layout of iShares' German Excel download: a sheet per topic, the table on 'Holdings'
    below 'Per <date>' in German month names, German column names, decimal commas, no ISIN and
    cash, futures and currency lines named in German."""
    cover = [["iShares Core MSCI Europe UCITS ETF EUR (Acc)"], ["NAV per 05.Okt.2026"]]
    file = suggest(workbook({"Fund Header": cover, "Holdings": GERMAN_ISHARES}))
    assert file.mapping.sheet == "Holdings" and file.as_of == date(2026, 10, 5)
    read = read_holdings(file)
    assert [(c.name, c.weight_pct, c.ticker, c.sector, c.currency) for c in read.constituents] == [
        ("ASML HOLDING", D("5.32"), "ASML", "IT", "EUR"),
        ("HSBC HOLDINGS PLC", D("94.25"), "HSBA", "Finanzwesen", "GBP"),
    ]
    assert read.covered_pct == D("99.57") and read.errors == []
    assert [d.split(":")[0] for d in read.dropped] == [
        "EUR CASH",
        "MSCI EUROPE INDEX DEC 26",
        "GBP/EUR",
    ]


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("05.Okt.2026", date(2026, 10, 5)),
        ("30.Sept.2026", date(2026, 9, 30)),
        ("1. M\u00e4rz 2026", date(2026, 3, 1)),
        ("31.Dez.2025", date(2025, 12, 31)),
        ("02 mei 2026", date(2026, 5, 2)),
    ],
)
def test_dates_with_german_and_dutch_month_names_are_read(text: str, expected: date) -> None:
    rows = [["Per", text], ["Name", "Weight (%)"], ["A", 100]]
    assert suggest(workbook({"S": rows})).as_of == expected
