"""ETF holdings for look-through (FR-MD-09): uploads with a preview, a monthly refresh from an
issuer's address or EODHD fundamentals, snapshots, and the stale notice."""

import json
from collections.abc import Callable
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path

import httpx
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select

from folio.config import Settings
from folio.db.engine import make_engine, make_session_factory
from folio.db.models_insight import EtfConstituent, EtfSnapshot
from folio.db.models_ledger import JobRequest, ProviderCall
from folio.db.models_strategy import Notification
from folio.jobs.context import JobContext
from folio.jobs.lookthrough import lookthrough_job
from folio.lookthrough.sources import parse_eodhd_holdings
from folio.marketdata.base import ProviderError
from folio.marketdata.budget import UsageTracker
from folio.marketdata.eodhd import EodhdProvider
from folio.marketdata.fallback import ProviderChain
from folio.security.secrets import SecretStore
from folio.settings_schema import HoldingsSource, LookThroughSettings
from folio.settings_store import save_section
from tests.conftest import PASSWORD, TEST_SECRET, USERNAME
from tests.marketdata_helpers import Scripted, client, make_listing, respond

D = Decimal
ISHARES = Path(__file__).resolve().parents[1] / "fixtures" / "lookthrough" / "ishares_style.csv"
NOW = datetime(2026, 10, 5, 6, tzinfo=UTC)


@pytest.fixture
def db(settings: Settings):  # type: ignore[no-untyped-def]
    with make_session_factory(make_engine(settings.db_url))() as session:
        yield session


@pytest.fixture
def api(make_client: Callable[..., TestClient], owner: None) -> TestClient:
    client_ = make_client()
    client_.post("/api/v1/auth/login", json={"username": USERNAME, "password": PASSWORD})
    return client_


def etf(db) -> int:  # type: ignore[no-untyped-def]
    instrument, _ = make_listing(db, symbols={"eodhd": "SXR8.XETRA"})
    db.commit()
    return instrument.id


def upload_to(api: TestClient, path: str, data: bytes, **form: str) -> httpx.Response:
    return api.post(path, files={"file": ("holdings.csv", data, "text/csv")}, data=form)


# --- upload, preview, snapshots -----------------------------------------------------------------


def test_a_file_is_previewed_then_stored_as_a_dated_snapshot(api: TestClient, db) -> None:  # type: ignore[no-untyped-def]
    instrument_id = etf(db)
    base = f"/api/v1/instruments/{instrument_id}/holdings"
    data = ISHARES.read_bytes()

    shown = upload_to(api, f"{base}/preview", data).json()
    assert (shown["holdings"], shown["as_of"], shown["covered_pct"]) == (3, "2026-10-03", "99.50")
    assert shown["top"][0]["name"] == "ALPHA TECH INC" and shown["errors"] == []
    assert [d.split(":")[0] for d in shown["dropped"]] == ["USD CASH", "FUTURES USD"]
    assert {h["label"] for h in shown["headers"]} >= {"Name", "Weight (%)"}
    assert api.get(base).json()["snapshots"] == []  # a preview stores nothing

    stored = upload_to(api, base, data)
    assert stored.status_code == 201
    snapshot = stored.json()["snapshot"]
    assert (snapshot["as_of"], snapshot["source"], snapshot["holdings"]) == ("2026-10-03", "csv", 3)

    # the same file again replaces that day's snapshot instead of doubling it
    upload_to(api, base, data)
    view = api.get(base).json()
    assert len(view["snapshots"]) == 1
    assert [c["name"] for c in view["top"]] == [
        "ALPHA TECH INC",
        "BETA BANK PLC",
        "CHIP MAKER NV",
    ]
    assert db.scalar(select(EtfConstituent.id).limit(1)) is not None
    assert len(db.scalars(select(EtfConstituent)).all()) == 3

    gone = api.delete(f"{base}/{view['snapshots'][0]['id']}")
    assert gone.status_code == 204
    assert api.get(base).json()["snapshots"] == []
    assert db.scalars(select(EtfConstituent)).all() == []


def excel_file() -> bytes:
    import io

    from openpyxl import Workbook

    book = Workbook()
    cover = book.active
    cover.title = "Cover"
    cover.append(["Example ETF factsheet"])
    sheet = book.create_sheet("Holdings")
    sheet.append(["Holdings as of", "02/10/2026"])
    sheet.append(["Weight (%)", "Name", "ISIN"])
    sheet.append([60, "ALPHA TECH INC", "US0000000001"])
    sheet.append([40, "BETA BANK PLC", "GB0000000002"])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def test_an_excel_workbook_is_previewed_with_its_sheets_then_stored(api: TestClient, db) -> None:  # type: ignore[no-untyped-def]
    base = f"/api/v1/instruments/{etf(db)}/holdings"
    data = excel_file()

    shown = upload_to(api, f"{base}/preview", data).json()
    assert [(s["name"], s["holdings"]) for s in shown["sheets"]] == [("Cover", 0), ("Holdings", 2)]
    assert shown["mapping"]["sheet"] == "Holdings"
    assert (shown["holdings"], shown["covered_pct"], shown["as_of"]) == (2, "100", "2026-10-02")
    assert [t["name"] for t in shown["top"]] == ["ALPHA TECH INC", "BETA BANK PLC"]

    # the owner can pick the other sheet: it has no holdings, and the preview says so
    cover = upload_to(api, f"{base}/preview", data, sheet="Cover")
    assert cover.status_code == 422 and "No header row" in cover.json()["detail"]
    again = upload_to(api, f"{base}/preview", data, sheet="Holdings").json()
    assert again["mapping"]["sheet"] == "Holdings" and again["holdings"] == 2

    stored = upload_to(api, base, data, mapping=json.dumps(shown["mapping"]))
    assert stored.status_code == 201 and stored.json()["snapshot"]["holdings"] == 2


def pdf_file() -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_font("Helvetica", size=10)
    pdf.add_page()
    pdf.cell(0, 8, "Holdings as of 02/10/2026", new_x="LMARGIN", new_y="NEXT")
    with pdf.table(col_widths=(90, 40)) as table:
        for cells in (
            ["Name", "Weight (%)"],
            ["ALPHA TECH INC", "60.00"],
            ["BETA BANK PLC", "40.00"],
        ):
            row = table.row()
            for cell in cells:
                row.cell(cell)
    return bytes(pdf.output())


def test_a_pdf_table_is_previewed_and_stored_like_any_other_file(api: TestClient, db) -> None:  # type: ignore[no-untyped-def]
    base = f"/api/v1/instruments/{etf(db)}/holdings"
    shown = upload_to(api, f"{base}/preview", pdf_file()).json()
    assert (shown["holdings"], shown["covered_pct"], shown["as_of"]) == (2, "100.00", "2026-10-02")
    assert shown["sheets"] == [] and shown["mapping"]["sheet"] is None
    stored = upload_to(api, base, pdf_file(), mapping=json.dumps(shown["mapping"]))
    assert stored.status_code == 201 and stored.json()["snapshot"]["holdings"] == 2
    scan = upload_to(api, f"{base}/preview", b"%PDF-1.4\nnot a pdf")
    assert scan.status_code == 422 and "could not be read" in scan.json()["detail"]


def test_the_job_refreshes_holdings_from_a_pdf_download(settings: Settings, db) -> None:  # type: ignore[no-untyped-def]
    instrument_id = etf(db)
    set_source(db, settings, instrument_id, HoldingsSource(url="https://issuer.example/h.pdf"))
    issuer = Scripted(lambda r: httpx.Response(200, content=pdf_file()))
    result = lookthrough_job(job_ctx(settings, issuer=issuer))
    assert result.status == "ok" and "2 holdings from url" in result.log


def test_the_job_refreshes_holdings_from_an_excel_download(settings: Settings, db) -> None:  # type: ignore[no-untyped-def]
    instrument_id = etf(db)
    set_source(
        db, settings, instrument_id, HoldingsSource(url="https://issuer.example/holdings")
    )  # no file extension: the kind is read from the content
    issuer = Scripted(lambda r: httpx.Response(200, content=excel_file()))
    result = lookthrough_job(job_ctx(settings, issuer=issuer))
    assert result.status == "ok" and "2 holdings from url" in result.log
    assert db.scalar(select(EtfSnapshot)).as_of == date(2026, 10, 2)


def test_a_file_that_is_not_a_holdings_table_is_explained(api: TestClient, db) -> None:  # type: ignore[no-untyped-def]
    base = f"/api/v1/instruments/{etf(db)}/holdings"
    refused = upload_to(api, base, b"Date,Amount\n2026-01-01,5\n")
    assert refused.status_code == 422
    assert "name column and a weight column" in refused.json()["detail"]
    unreadable = upload_to(api, base, b"Name,Weight\nAlpha,80\nBeta,70\n")
    assert unreadable.status_code == 422 and "150.0 %" in unreadable.json()["detail"]
    other = upload_to(api, "/api/v1/instruments/9999/holdings", b"Name,Weight\nA,100\n")
    assert other.status_code == 404


def test_the_column_match_can_be_changed_on_the_preview(api: TestClient, db) -> None:  # type: ignore[no-untyped-def]
    base = f"/api/v1/instruments/{etf(db)}/holdings"
    data = b"Name,Weight,Share\nAlpha,10,60\nBeta,20,40\n"
    first = upload_to(api, f"{base}/preview", data).json()
    assert first["top"][0]["name"] == "Beta"  # matched "Weight", as written
    mapping = {**first["mapping"], "weight": 2}
    again = upload_to(api, f"{base}/preview", data, mapping=json.dumps(mapping))
    assert [t["weight_pct"] for t in again.json()["top"]] == ["60", "40"]


def test_a_snapshot_belongs_to_its_own_etf(api: TestClient, db) -> None:  # type: ignore[no-untyped-def]
    first = etf(db)
    base = f"/api/v1/instruments/{first}/holdings"
    snapshot_id = upload_to(api, base, ISHARES.read_bytes()).json()["snapshot"]["id"]
    second, _ = make_listing(db, ticker="VUAA", isin="IE00BFMXXD54")
    db.commit()
    wrong = api.delete(f"/api/v1/instruments/{second.id}/holdings/{snapshot_id}")
    assert wrong.status_code == 404


# --- the monthly job ----------------------------------------------------------------------------


def job_ctx(
    settings: Settings,
    issuer: Scripted | None = None,
    eodhd: Scripted | None = None,
    now: datetime = NOW,
) -> JobContext:
    factory = make_session_factory(make_engine(settings.db_url))
    usage = UsageTracker(factory, lambda p: 20 if p == "eodhd" else 0)
    return JobContext(
        session_factory=factory,
        chain_for=lambda db: ProviderChain([]),
        ecb_for=lambda db: None,  # type: ignore[arg-type, return-value]
        now=lambda: now,
        issuer_for=lambda db: None if issuer is None else client("issuer", issuer),
        eodhd_for=lambda db: (
            None
            if eodhd is None
            else EodhdProvider(client("eodhd", eodhd, usage=usage), "fake-eodhd-key-1234")
        ),
    )


def set_source(db, settings: Settings, instrument_id: int, source: HoldingsSource) -> None:  # type: ignore[no-untyped-def]
    store = SecretStore(db, TEST_SECRET)
    save_section(
        db, store, "lookthrough", LookThroughSettings(sources={str(instrument_id): source})
    )
    db.commit()


def test_the_job_refreshes_holdings_from_the_issuers_address(settings: Settings, db) -> None:  # type: ignore[no-untyped-def]
    instrument_id = etf(db)
    set_source(
        db, settings, instrument_id, HoldingsSource(url="https://issuer.example/holdings.csv")
    )
    issuer = Scripted(lambda r: httpx.Response(200, content=ISHARES.read_bytes()))
    result = lookthrough_job(job_ctx(settings, issuer=issuer))
    assert result.status == "ok" and "3 holdings from url" in result.log
    snapshot = db.scalar(select(EtfSnapshot))
    assert (snapshot.source, snapshot.as_of) == ("url", date(2026, 10, 3))  # the file's own date
    assert str(issuer.requests[0].url) == "https://issuer.example/holdings.csv"


def test_a_web_page_instead_of_a_file_is_reported_not_parsed(settings: Settings, db) -> None:  # type: ignore[no-untyped-def]
    instrument_id = etf(db)
    set_source(db, settings, instrument_id, HoldingsSource(url="https://issuer.example/page"))
    page = Scripted(lambda r: httpx.Response(200, content=b"<!DOCTYPE html><html></html>"))
    result = lookthrough_job(job_ctx(settings, issuer=page))
    assert result.status == "failed" and "returned a web page, not a holdings file" in result.log
    assert db.scalar(select(EtfSnapshot)) is None


def test_eodhd_fundamentals_are_read_and_cost_ten_calls(settings: Settings, db) -> None:  # type: ignore[no-untyped-def]
    instrument_id = etf(db)
    set_source(db, settings, instrument_id, HoldingsSource(eodhd=True))
    answer = Scripted(lambda r: respond("eodhd_fundamentals_etf.json"))
    result = lookthrough_job(job_ctx(settings, eodhd=answer))
    assert result.status == "ok" and "3 holdings from eodhd" in result.log
    assert answer.requests[0].url.path == "/api/fundamentals/SXR8.XETRA"
    assert db.scalar(select(ProviderCall.count).where(ProviderCall.provider == "eodhd")) == 10
    rows = db.scalars(select(EtfConstituent).order_by(EtfConstituent.weight_pct.desc())).all()
    assert [(r.name, r.weight_pct, r.country) for r in rows][0] == (
        "Alpha Tech Inc",
        D("40.25"),
        "USA",
    )


def test_a_refusal_for_fundamentals_names_the_plan_not_just_the_key(settings: Settings, db) -> None:  # type: ignore[no-untyped-def]
    """A free EODHD key is refused (HTTP 403) for fund holdings: the message says why and what to
    do instead, and it is not the generic "check the API key"."""
    instrument_id = etf(db)
    set_source(db, settings, instrument_id, HoldingsSource(eodhd=True))
    refused = Scripted(lambda r: httpx.Response(403, json={"message": "Forbidden"}))
    result = lookthrough_job(job_ctx(settings, eodhd=refused))
    assert result.status == "failed"
    assert "HTTP 403" in result.log and "Fundamentals data" in result.log
    assert "free plan does not include" in result.log and "Add a holdings file" in result.log
    assert "Check the API key" not in result.log


def test_eodhd_without_a_key_says_what_to_do(settings: Settings, db) -> None:  # type: ignore[no-untyped-def]
    set_source(db, settings, etf(db), HoldingsSource(eodhd=True))
    result = lookthrough_job(job_ctx(settings, eodhd=None))
    assert result.status == "failed" and "EODHD needs an API key" in result.log


def test_an_old_snapshot_gets_one_notice_a_month(settings: Settings, db) -> None:  # type: ignore[no-untyped-def]
    instrument_id = etf(db)
    base = {"instrument_id": instrument_id, "source": "csv", "covered_pct": D(100)}
    db.add(EtfSnapshot(as_of=date(2026, 8, 1), **base))  # 65 days before NOW
    db.commit()
    lookthrough_job(job_ctx(settings))
    notices = db.scalars(select(Notification).where(Notification.source == "system")).all()
    assert len(notices) == 1 and "65 days old" in notices[0].title
    assert notices[0].severity == "low" and notices[0].link == f"/holdings/{instrument_id}"
    lookthrough_job(job_ctx(settings))  # once a month, not every run
    assert len(db.scalars(select(Notification).where(Notification.source == "system")).all()) == 1


def test_the_refresh_button_queues_the_job_and_the_source_is_saved(api: TestClient, db) -> None:  # type: ignore[no-untyped-def]
    instrument_id = etf(db)
    base = f"/api/v1/instruments/{instrument_id}/holdings"
    saved = api.put(f"{base}/source", json={"url": "https://issuer.example/h.csv", "eodhd": False})
    assert saved.json() == {"url": "https://issuer.example/h.csv", "eodhd": False}
    assert api.get(base).json()["source"]["url"] == "https://issuer.example/h.csv"
    assert api.post(f"{base}/refresh").status_code == 202
    queued = db.scalars(select(JobRequest).where(JobRequest.job == "lookthrough")).all()
    assert [q.params for q in queued] == [{"instrument_id": instrument_id}]
    cleared = api.put(f"{base}/source", json={"url": None, "eodhd": False})
    assert cleared.json() == {"url": None, "eodhd": False}


# --- sources, in isolation -----------------------------------------------------------------------


def test_eodhd_with_no_holdings_names_the_plan_it_needs() -> None:
    read = parse_eodhd_holdings({"General": {"Code": "X"}})
    assert read.errors and "Fundamentals plan" in read.errors[0]


def test_a_non_http_address_is_refused_before_any_request() -> None:
    from folio.lookthrough.sources import fetch_issuer_file

    seen = Scripted(lambda r: httpx.Response(200, content=b""))
    with pytest.raises(ProviderError, match="http"):
        fetch_issuer_file(client("issuer", seen), "file:///etc/passwd")
    assert seen.requests == []
